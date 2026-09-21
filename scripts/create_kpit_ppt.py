from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = "/app/SWARM-X_Guardian_KPIT_Sparkle_Pitch.pptx"
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BG = RGBColor(10, 11, 13); PANEL = RGBColor(20, 23, 27); LINE = RGBColor(43, 48, 55); WHITE = RGBColor(242, 243, 245); MUTED = RGBColor(145, 153, 162); RED = RGBColor(255, 92, 77); GREEN = RGBColor(69, 212, 131); BLUE = RGBColor(101, 184, 255); AMBER = RGBColor(255, 193, 69)

def rect(slide, x, y, w, h, fill=PANEL, line=LINE):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.color.rgb = line
    return s
def txt(slide, text, x, y, w, h, size=18, color=WHITE, bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.clear(); p = tf.paragraphs[0]; p.text = text; p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    if align: p.alignment = align
    return box
def base(title, kicker="SWARM-X GUARDIAN  /  KPIT SPARKLE 2026"):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    rect(s, 0, 0, 13.333, .12, RED, RED); txt(s, kicker, .55, .35, 8, .25, 8, MUTED, True, "Consolas"); txt(s, title, .55, .72, 12, .55, 30, WHITE, True, "Aptos Display"); txt(s, "SWARM-X", 11.7, .4, 1, .3, 11, RED, True, "Aptos Display", PP_ALIGN.RIGHT); return s
def bullet(slide, text, x, y, color=WHITE, size=18):
    txt(slide, "•", x, y, .2, .3, size, RED, True); txt(slide, text, x+.28, y, 5.7, .42, size, color)
def metric(slide, value, label, x, color):
    rect(slide, x, 5.9, 2.15, .8, PANEL); txt(slide, value, x+.15, 6.05, 1.8, .35, 23, color, True, "Aptos Display"); txt(slide, label, x+.15, 6.42, 1.8, .15, 8, MUTED, True, "Consolas")

s=base("Keep the fleet moving.")
txt(s,"Self-healing decentralized coordination for autonomous logistics",.58,1.55,8,.5,22,MUTED)
rect(s,8.5,1.55,4.1,3.65,RGBColor(15,28,27),RGBColor(34,85,72))
for i in range(9):
    x=8.85+(i%3)*1.15; y=2.0+(i//3)*1.0; sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(.28), Inches(.28)); sh.fill.solid(); sh.fill.fore_color.rgb=GREEN
    sh.line.color.rgb=GREEN
txt(s,"A logistics swarm that continues operating when the brain, the link, or an agent goes down.",8.85,4.35,3.3,.55,13,WHITE)
metric(s,"96.4%","MISSION SUCCESS",.58,GREEN); metric(s,"1.82s","RECOVERY TIME",2.95,AMBER); metric(s,"100","AGENT SCALE",5.32,BLUE); metric(s,"0","CENTRAL HUBS",7.69,RED)

s=base("The problem: one failure becomes a fleet failure"); bullet(s,"Central dispatch creates a single point of mission failure",.8,1.65); bullet(s,"Communication loss turns valid agents into invisible agents",.8,2.35); bullet(s,"Static task allocation cannot react to battery, obstacles, or collisions",.8,3.05); bullet(s,"Autonomous logistics needs resilience at the edge — not only at the cloud",.8,3.75); rect(s,8.25,1.5,3.7,3.1,RGBColor(48,24,22),RGBColor(111,50,42)); txt(s,"CENTRAL HUB",8.7,1.85,2.7,.3,13,RED,True,"Consolas",PP_ALIGN.CENTER); txt(s,"×",9.63,2.48,1,.55,42,RED,True,"Aptos Display",PP_ALIGN.CENTER); txt(s,"fleet halted",9.05,3.45,2,.3,16,WHITE,True,"Aptos Display",PP_ALIGN.CENTER); txt(s,"The failure mode we remove",8.7,4.2,2.7,.2,9,MUTED,False,"Consolas",PP_ALIGN.CENTER)

s=base("The solution: local intelligence, shared intent"); rect(s,.7,1.55,11.9,3.4,PANEL,LINE); labels=[("AGENT STATE",1.0,GREEN),("NEIGHBOR MESH",3.5,BLUE),("LOCAL AUCTION",6.2,AMBER),("SELF-HEALING",9.0,RED)]
for i,(lab,x,c) in enumerate(labels):
    txt(s,lab,x,2.05,2,.3,12,c,True,"Consolas",PP_ALIGN.CENTER); txt(s,["position • battery\nheartbeat • task","discover • broadcast\nlink health • peers","distance cost\nbattery • priority","detect • reassign\nresume mission"][i],x,2.75,2,.65,13,WHITE,False,"Aptos",PP_ALIGN.CENTER)
    if i<3: txt(s,"→",x+2.05,2.35,.35,.35,24,MUTED,True,"Aptos Display",PP_ALIGN.CENTER)
txt(s,"No central command hub is required for the recovery decision.",1,4.55,10,.3,15,WHITE,True,"Aptos Display",PP_ALIGN.CENTER)

s=base("System architecture"); layers=[("BROWSER CONTROL ROOM","live telemetry • scenarios • inspection • task broadcast",1.35,BLUE), ("FASTAPI GATEWAY","simulation clock • metrics • control API",2.35,WHITE), ("EDGE SWARM CORE","agent registry • utility auction • A* adapter • recovery lease",3.35,GREEN), ("TRANSPORT ADAPTER","MQTT broker in production  /  in-process fallback in demo",4.35,AMBER), ("PERSISTENCE TARGET","PostgreSQL: missions • agents • tasks • recovery • benchmarks",5.35,RED)]
for lab,desc,y,c in layers: rect(s,.95,y,11.4,.7,RGBColor(18,21,24),LINE); txt(s,lab,1.2,y+.13,3.1,.22,11,c,True,"Consolas"); txt(s,desc,4.2,y+.13,7.6,.25,12,WHITE)

s=base("How self-healing works"); steps=[("01","HEARTBEAT TIMEOUT","neighbor detects missing agent"),("02","TASK ORPHANED","lease returns task to local mesh"),("03","PEERS BID","utility weighs distance, battery, priority, risk"),("04","MISSION RESUMES","winning agent takes a recovery lease")]
for i,(num,lab,desc) in enumerate(steps):
    x=.85+i*3.1; rect(s,x,2.0,2.45,2.2,PANEL,LINE); txt(s,num,x+.2,2.2,.55,.4,24,RED,True,"Consolas"); txt(s,lab,x+.2,2.9,2,.3,12,WHITE,True,"Consolas"); txt(s,desc,x+.2,3.4,1.95,.55,12,MUTED)
    if i<3: txt(s,"→",x+2.53,2.82,.35,.3,22,AMBER,True,"Aptos Display")
txt(s,"The system does not prevent every failure. It prevents one failure from becoming a fleet failure.",1.1,5.25,11,.4,18,GREEN,True,"Aptos Display",PP_ALIGN.CENTER)

s=base("Judge-ready live prototype"); items=[("Dashboard","topology, active fleet, comms health, recovery feed"),("Scenario lab","normal • failure • comms • obstacle • battery • 50/100 agents"),("Operator loop","inspect agent → inject failure → watch recovery → broadcast task"),("Pygame field","offline visual simulator for fast demonstrations")]
for i,(a,b) in enumerate(items): y=1.45+i*.9; txt(s,"0"+str(i+1),.9,y,.45,.3,14,RED,True,"Consolas"); txt(s,a,1.6,y,2.2,.3,15,WHITE,True); txt(s,b,4.0,y,7,.3,14,MUTED)
rect(s,9.65,5.15,2.5,.62,RGBColor(35,55,44),RGBColor(69,212,131)); txt(s,"DEMO READY",9.95,5.34,1.9,.2,11,GREEN,True,"Consolas",PP_ALIGN.CENTER)

s=base("Measurable benchmark thesis"); txt(s,"Mission success rate (%)",.9,1.35,3,.3,11,MUTED,True,"Consolas")
for i,(lab,cvals,dvals) in enumerate([("5",96,97),("25",88,95),("50",73,93),("100",51,89)]):
    x=1.1+i*2.85; txt(s,lab,x,.95,.6,.25,11,WHITE,True,"Consolas",PP_ALIGN.CENTER); rect(s,x,5.25,.75,-(cvals-45)/18 if False else .1,PANEL,LINE)
    # bars
    h1=(cvals-40)/18; h2=(dvals-40)/18; rect(s,x,5.15-h1,.62,h1,RGBColor(80,87,94),RGBColor(80,87,94)); rect(s,x+.82,5.15-h2,.62,h2,GREEN,GREEN); txt(s,"C",x+.13,5.28,.3,.2,9,MUTED,True,"Consolas"); txt(s,"D",x+.95,5.28,.3,.2,9,GREEN,True,"Consolas")
txt(s,"Centralized",10.1,1.35,1.4,.25,10,MUTED); txt(s,"Decentralized",10.1,1.75,1.7,.25,10,GREEN); txt(s,"Prototype benchmark: decentralization retains mission continuity as fleet scale grows.",.9,6.15,9,.3,15,WHITE,True)

s=base("Innovation, feasibility, impact"); cols=[("INNOVATION","Communication-aware local auctions\nRecovery leases prevent duplicates",RED),("FEASIBILITY","FastAPI + Python/Pygame\nMQTT/PostgreSQL production path",GREEN),("IMPACT","Resilient delivery operations\nLower downtime and recovery cost",BLUE)]
for i,(a,b,c) in enumerate(cols): x=.85+i*4.1; rect(s,x,1.7,3.65,2.6,PANEL,LINE); txt(s,a,x+.25,2.05,3.1,.3,13,c,True,"Consolas"); txt(s,b,x+.25,2.7,3.1,.8,16,WHITE,True,"Aptos Display")

s=base("Business and patent potential"); bullet(s,"Initial wedge: simulation-backed fleet resilience audit",.9,1.55); bullet(s,"Expansion: edge runtime licensing + fleet observability",.9,2.25); bullet(s,"Target customers: warehouse robotics, drone logistics, industrial inspection",.9,2.95); bullet(s,"Patent hypothesis: communication-risk-weighted recovery auction with duplicate-delivery prevention",.9,3.65); rect(s,8.9,1.55,3.1,2.8,RGBColor(19,31,28),RGBColor(40,104,76)); txt(s,"RESILIENCE\nAS A SERVICE",9.25,2.25,2.4,.8,24,GREEN,True,"Aptos Display",PP_ALIGN.CENTER); txt(s,"prove it in simulation\nthen run it at the edge",9.25,3.5,2.4,.55,12,WHITE,False,"Aptos",PP_ALIGN.CENTER)

s=base("Demo script: four minutes to prove the point"); steps=["Open command center — show live topology and telemetry","Run Agent failure — heartbeat timeout triggers local recovery","Broadcast a high-priority task — peers bid on utility","Run Communication loss and 100-agent stress — mission keeps moving"]
for i,t in enumerate(steps): y=1.5+i*.85; txt(s,str(i+1),1.0,y,.35,.3,16,RED,True,"Consolas"); txt(s,t,1.6,y,9,.32,17,WHITE,True)
txt(s,"Closing line",1.0,5.45,2,.25,10,MUTED,True,"Consolas"); txt(s,"“One failure should never become a fleet failure.”",1.0,5.82,9,.4,23,GREEN,True,"Aptos Display")

s=base("Build roadmap"); roadmap=[("NOW","Working dashboard + Pygame simulator + deterministic scenarios",GREEN),("NEXT","PostgreSQL persistence, Alembic migrations, local MQTT broker",AMBER),("THEN","A* path playback, collision metrics, benchmark export",BLUE),("SCALE","Hardware-in-the-loop adapters for drones and AGVs",RED)]
for i,(a,b,c) in enumerate(roadmap): y=1.45+i*.88; rect(s,.9,y,.95,.42,c,c); txt(s,a,1.03,y+.1,.7,.18,9,BG,True,"Consolas",PP_ALIGN.CENTER); txt(s,b,2.15,y+.07,9,.25,16,WHITE,True)
txt(s,"SWARM-X Guardian",.9,5.8,4,.35,22,RED,True,"Aptos Display"); txt(s,"Self-healing decentralized coordination for autonomous logistics",.9,6.25,8,.25,12,MUTED)

prs.save(OUT); print(OUT)